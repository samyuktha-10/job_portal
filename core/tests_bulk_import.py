"""Bulk Excel candidate import: creates real User + JobSeekerProfile rows,
skips duplicates, reports row errors, gates access by role."""
import io

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .bulk_import_views import DEFAULT_PASSWORD
from .models import JobSeekerProfile, Profile


def _xlsx(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ['Full Name', 'Email', 'Phone', 'Location', 'Education',
                             'Skills', 'Experience', 'Preferred Job Type', 'Certificates'])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        'candidates.xlsx', buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class BulkCandidateImportTests(TestCase):
    def setUp(self):
        self.employer = User.objects.create_user('employer', 'e@x.com', 'pass1234')
        Profile.objects.create(user=self.employer, is_employer=True, company_name='Co')
        self.seeker = User.objects.create_user('seeker', 's@x.com', 'pass1234')
        JobSeekerProfile.objects.create(user=self.seeker, full_name='Plain Seeker',
                                        phone='9999999999')
        self.admin = User.objects.create_superuser('root', 'r@x.com', 'pass1234')

    def _post(self, upload, url='bulk_candidate_import'):
        return self.client.post(reverse(url), {'excel_file': upload})

    def test_employer_page_renders_and_job_seeker_blocked(self):
        self.client.force_login(self.employer)
        response = self.client.get(reverse('bulk_candidate_import'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Bulk Candidate Upload')

        self.client.force_login(self.seeker)
        response = self.client.get(reverse('bulk_candidate_import'))
        self.assertEqual(response.status_code, 302)  # employer_required turns them away

    def test_super_admin_url_works(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('admin_bulk_candidate_import'))
        self.assertEqual(response.status_code, 200)

    def test_upload_creates_real_profiles_with_login(self):
        self.client.force_login(self.employer)
        upload = _xlsx([
            ['Anu Sharma', 'anu@example.com', '9876543210', 'Chennai', 'B.E. CSE',
             'Python, Django', '1-3', 'full-time', 'AWS CPA'],
            ['Karthik R', '', '9000000001', 'Madurai', 'MCA', 'Java', 'fresher',
             'internship', ''],
        ])
        response = self._post(upload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['created'], 2)

        anu = User.objects.get(email='anu@example.com')
        profile = anu.jobseeker_profile
        self.assertEqual(profile.full_name, 'Anu Sharma')
        self.assertEqual(profile.phone, '9876543210')
        self.assertEqual(profile.skills, 'Python, Django')
        self.assertEqual(profile.preferred_job_type, 'full-time')
        self.assertTrue(profile.is_experienced)
        self.assertFalse(profile.whatsapp_opted_in)

        karthik = JobSeekerProfile.objects.get(full_name='Karthik R')
        self.assertFalse(karthik.is_experienced)
        self.assertTrue(karthik.user.email.endswith('@imported.candidates.deploynix'))
        self.assertEqual(response.context['synthesized'], 1)

        # the created account really can log in with the default password
        self.assertTrue(self.client.login(username=anu.username,
                                          password=DEFAULT_PASSWORD))

    def test_duplicates_skipped_and_errors_reported(self):
        self.client.force_login(self.employer)
        upload = _xlsx([
            ['Plain Seeker', 's@x.com', '9999999999', '', '', '', '', '', ''],
            ['', 'nobody@example.com', '', '', '', '', '', '', ''],
            ['Valid Person', 'valid@example.com', '', '', '', '', '', '', ''],
        ])
        response = self._post(upload)
        self.assertEqual(response.context['created'], 1)
        self.assertEqual(len(response.context['duplicates']), 1)
        self.assertEqual(response.context['duplicates'][0][1], 's@x.com')
        self.assertEqual(len(response.context['row_errors']), 1)
        self.assertIn('Missing Full Name', response.context['row_errors'][0][1])
        self.assertFalse(User.objects.filter(email='nobody@example.com').exists())

    def test_header_below_title_row_and_fuzzy_names(self):
        self.client.force_login(self.employer)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Deploynix Placement Drive - Batch 2026'])
        sheet.append([''])
        sheet.append(['S.No', 'Name of the Candidate', 'Mail', 'Mobile No',
                      'City', 'Qualification', 'Key Skills', 'Years'])
        sheet.append([1, 'Ravi Teja', 'ravi@example.com', '9812345678', 'Hyderabad',
                      'B.Tech', 'Python, SQL', '1-3'])
        sheet.append([2, 'Meena K', '', '9812345679', 'Chennai', 'MCA', 'Java', 'fresher'])
        buffer = io.BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'drive.xlsx', buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response = self._post(upload)
        self.assertEqual(response.status_code, 200)
        ravi = JobSeekerProfile.objects.get(full_name='Ravi Teja')
        self.assertEqual(ravi.user.email, 'ravi@example.com')
        self.assertEqual(ravi.phone, '9812345678')
        self.assertEqual(ravi.location, 'Hyderabad')
        self.assertEqual(ravi.skills, 'Python, SQL')
        self.assertTrue(JobSeekerProfile.objects.filter(full_name='Meena K').exists())

    def test_unknown_headers_error_lists_columns(self):
        self.client.force_login(self.employer)
        upload = _xlsx([['X', 'Y'], ], headers=['Column A', 'Column B'])
        response = self._post(upload)
        self.assertContains(response, 'Columns in your file')
        self.assertContains(response, 'Column A')
        self.assertEqual(JobSeekerProfile.objects.count(), 1)  # only setUp's

    def test_email_less_rows_deduped_by_phone_on_reupload(self):
        self.client.force_login(self.employer)
        rows = [['No Email Person', '', '9000000007', '', '', '', '', '', '']]
        self._post(_xlsx(rows))
        self._post(_xlsx(rows))  # same file again
        self.assertEqual(
            JobSeekerProfile.objects.filter(full_name='No Email Person').count(), 1)

    def test_non_xlsx_rejected(self):
        self.client.force_login(self.employer)
        upload = SimpleUploadedFile('candidates.csv', b'a,b\n1,2\n', content_type='text/csv')
        response = self._post(upload)
        self.assertEqual(response.context.get('created'), None)
        self.assertContains(response, '.xlsx')
        self.assertEqual(JobSeekerProfile.objects.count(), 1)  # only setUp's

    def test_template_download_is_workbook(self):
        self.client.force_login(self.employer)
        response = self.client.get(reverse('bulk_import_template'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml.sheet', response['Content-Type'])
        workbook = load_workbook(io.BytesIO(b''.join(response.streaming_content
                                                     if response.streaming
                                                     else [response.content])))
        self.assertEqual(workbook.active['A1'].value, 'Full Name')
