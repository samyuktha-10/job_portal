"""Bulk candidate-profile import from an Excel (.xlsx) workbook.

Uploading a spreadsheet of candidates creates REAL platform accounts:
one ``User`` (login enabled with a default password) plus one
``JobSeekerProfile`` per row, so the candidates immediately appear in
candidate search / job-seeker management instead of sitting in a raw table.
Available to employers (Candidates menu) and super admins (Control Panel).
"""
import io
import re

from django.contrib import messages
from django.contrib.auth.hashers import make_password
from django.contrib.auth.models import User
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse

from openpyxl import Workbook, load_workbook

from .decorators import admin_required, employer_required
from .models import JobSeekerProfile

DEFAULT_PASSWORD = 'Deploynix@123'
MAX_ROWS = 5000

TEMPLATE_HEADERS = [
    'Full Name', 'Email', 'Phone', 'Location', 'Education',
    'Skills', 'Experience', 'Preferred Job Type', 'Certificates',
]

HEADER_ALIASES = {
    'full_name': ['full_name', 'name', 'candidate_name', 'student_name', 'applicant_name'],
    'email': ['email', 'email_id', 'mail', 'email_address'],
    'phone': ['phone', 'mobile', 'contact', 'phone_number', 'mobile_number', 'contact_number'],
    'location': ['location', 'city', 'place', 'address'],
    'education': ['education', 'qualification', 'degree', 'highest_qualification'],
    'skills': ['skills', 'skills_known', 'tech_skills', 'key_skills'],
    'experience': ['experience', 'experience_years', 'exp', 'years_of_experience'],
    'preferred_job_type': ['preferred_job_type', 'job_type', 'looking_for'],
    'certificates': ['certificates', 'certs', 'certifications'],
}

JOB_TYPES = {'full-time', 'part-time', 'internship', 'remote'}


def _norm_header(value):
    return re.sub(r'[^a-z0-9]+', '_', str(value or '').strip().lower()).strip('_')


def _slug(value):
    return re.sub(r'[^a-z0-9._-]+', '-', str(value or '').strip().lower()).strip('.-')


def _cell(row, index):
    if index is None or index >= len(row):
        return ''
    value = row[index]
    return '' if value is None else str(value).strip()


def _parse_workbook(upload):
    """Return (rows, error). rows = list of (excel_row_no, dict-of-fields)."""
    if not upload.name.lower().endswith('.xlsx'):
        return None, 'Please upload an Excel .xlsx file (open older .xls files in ' \
                     'Excel and Save As .xlsx first).'
    try:
        workbook = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
    except Exception:
        return None, 'Could not read that file as a valid .xlsx workbook.'

    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        return None, 'The workbook is empty - add a header row first.'

    positions = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            for i, cell in enumerate(header):
                if _norm_header(cell) == alias:
                    positions[field] = i
                    break
            if field in positions:
                break
    if 'full_name' not in positions:
        return None, 'No "Full Name" column found. Download the template to see the ' \
                     'expected header row.'

    rows, count = [], 0
    for number, row in enumerate(iterator, start=2):
        if all(cell in (None, '') or str(cell).strip() == '' for cell in row):
            continue
        count += 1
        if count > MAX_ROWS:
            return None, f'Too many rows - the upload limit is {MAX_ROWS} candidates ' \
                         'per file.'
        rows.append((number, {field: _cell(row, idx) for field, idx in positions.items()}))
    if not rows:
        return None, 'The workbook has a header but no data rows.'
    return rows, None


def _build_users_and_profiles(rows):
    """Validate rows; return (users, profiles_data, duplicates, errors, synthesized)."""
    existing_emails = {u['email'] for u in User.objects.exclude(email='').values('email')}
    existing_emails = {e.lower() for e in existing_emails}
    existing_usernames = set(User.objects.values_list('username', flat=True))
    seen_emails = set()

    users, profiles_data, duplicates, errors = [], [], [], []
    synthesized = 0

    for number, data in rows:
        full_name = data.get('full_name', '')
        if not full_name:
            errors.append((number, 'Missing Full Name - row skipped.'))
            continue

        email = data.get('email', '').lower()
        if email:
            if '@' not in email:
                errors.append((number, f'"{email}" is not a valid email - row skipped.'))
                continue
            if email in existing_emails or email in seen_emails:
                duplicates.append((number, email))
                continue
            seen_emails.add(email)

        phone = re.sub(r'[^0-9+]', '', data.get('phone', ''))[:15]
        experience = data.get('experience', '')[:100]
        job_type = _slug(data.get('preferred_job_type', '')).replace('_', '-')
        if job_type not in JOB_TYPES:
            job_type = ''

        base = _slug(email.split('@')[0] if email else full_name) or 'candidate'
        base = base[:120] or 'candidate'
        username, suffix = base, 1
        while username in existing_usernames:
            suffix += 1
            username = f'{base[:120]}{suffix}'
        existing_usernames.add(username)

        if not email:
            email = f'{username}@imported.candidates.deploynix'
            synthesized += 1

        parts = full_name.split(None, 1)
        user = User(
            username=username, email=email,
            first_name=parts[0][:150], last_name=(parts[1] if len(parts) > 1 else '')[:150],
        )
        users.append(user)
        profiles_data.append({
            'full_name': full_name[:200], 'phone': phone,
            'location': data.get('location', '')[:150],
            'education': data.get('education', '')[:200],
            'skills': data.get('skills', '')[:300],
            'experience': experience,
            'preferred_job_type': job_type,
            'certificates': data.get('certificates', ''),
            'is_experienced': bool(experience) and experience.lower() not in ('fresher', '0'),
        })
    return users, profiles_data, duplicates, errors, synthesized


def bulk_candidate_import(request):
    context = {
        'template_url': (reverse('admin_bulk_import_template') if request.user.is_superuser
                         else reverse('bulk_import_template')),
        'default_password': DEFAULT_PASSWORD,
    }

    if request.method == 'POST':
        upload = request.FILES.get('excel_file')
        if not upload:
            messages.error(request, 'Choose an .xlsx file to upload first.')
            return render(request, 'core/bulk_candidate_import.html', context)

        rows, error = _parse_workbook(upload)
        if error:
            messages.error(request, error)
            return render(request, 'core/bulk_candidate_import.html', context)

        users, profiles_data, duplicates, errors, synthesized = _build_users_and_profiles(rows)
        created = 0
        if users:
            hashed = make_password(DEFAULT_PASSWORD)  # hashed once, shared default
            for user in users:
                user.password = hashed
            with transaction.atomic():
                created_users = User.objects.bulk_create(users, batch_size=500)
                JobSeekerProfile.objects.bulk_create([
                    JobSeekerProfile(user=user, whatsapp_opted_in=False, **data)
                    for user, data in zip(created_users, profiles_data)
                ], batch_size=500)
            created = len(created_users)
        context.update({
            'created': created,
            'duplicates': duplicates,
            'row_errors': errors,
            'synthesized': synthesized,
            'total_rows': len(rows),
        })
        if created:
            messages.success(request, f'{created} candidate profiles created from the '
                                      'Excel file.')
        if not created and not duplicates:
            messages.error(request, 'No candidate profiles could be created - see the '
                                    'row report below.')
        return render(request, 'core/bulk_candidate_import.html', context)

    return render(request, 'core/bulk_candidate_import.html', context)


def bulk_import_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Candidates'
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(['Anu Sharma', 'anu.sharma@example.com', '9876543210', 'Chennai',
                  'B.E. Computer Science', 'Python, Django, SQL', '1-3', 'full-time',
                  'AWS Certified Cloud Practitioner'])
    sheet.append(['Karthik R', '', '9000000001', 'Madurai', 'MCA', 'Java, Spring',
                  'fresher', 'internship', ''])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="candidate_import_template.xlsx"'
    workbook.save(response)
    return response


employer_bulk_candidate_import = employer_required(bulk_candidate_import)
admin_bulk_candidate_import = admin_required(bulk_candidate_import)
employer_bulk_import_template = employer_required(bulk_import_template)
admin_bulk_import_template = admin_required(bulk_import_template)
